#!/usr/bin/env python3
"""
build_data.py · Convierte los archivos de programas (.xlsm de CURSOS_2026.zip) en los JSON que lee el sitio.

Uso:
    python herramientas/build_data.py     # lee entrada/*.xlsm  y escribe datos/horarios.json + datos/meta.json
    python herramientas/build_data.py --entrada X --salida Y

Requiere:  pip install openpyxl
"""
import argparse, glob, json, os, re, sys, datetime
import openpyxl

DIAS = {"LUN": "LUN", "MAR": "MAR", "MIE": "MIE", "JUE": "JUE", "VIE": "VIE", "SAB": "SAB", "DOM": "DOM"}
HDR = re.compile(r"^\s*(\w+)\s*-\s*\(\s*([\w-]+)\s*\)\s*(.*)$")
HORA = re.compile(r"(\d{1,2}):(\d{2})\s*(AM|PM)", re.I)


def minutos(t):
    h, m, ap = HORA.search(t).groups()
    h = int(h) % 12 + (12 if ap.upper() == "PM" else 0)
    return h * 60 + int(m)


def rango(txt):
    a, b = HORA.findall(txt)[:2]
    return minutos("%s:%s %s" % a), minutos("%s:%s %s" % b)


def celdas(row):
    return [c for c in row]


PERIODO = re.compile(r"PERIODO\s+(\d{6})")


def leer_programa(path):
    k = os.path.splitext(os.path.basename(path))[0].upper()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    nombres = {}
    cursos = []
    periodos = set()
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if ws.title.lower() == "principal":
            for r in rows:
                for c in r[:1]:
                    m = PERIODO.search(str(c or ""))
                    if m:
                        periodos.add(m.group(1))
                if r and r[0] and re.match(r"^[A-Z]{3,5}-\d+$", str(r[0]).strip()):
                    nombres[str(r[0]).strip()] = str(r[1] or "").strip()
            continue
        cod, mod = None, ""
        secs, cur, cols = [], None, {}
        for r in rows:
            a = r[0]
            if isinstance(a, str) and HDR.match(a) and cod is None and "(" in a and "-" in a and "PRESENCIAL" not in a.upper():
                cod = HDR.match(a).group(2)
                continue
            if isinstance(a, str) and "PRESENCIAL" in a.upper() and cod and not secs and cur is None:
                mod = a.strip()
                continue
            if a == "NRC:":
                cols = {str(v).rstrip(":").strip(): i for i, v in enumerate(r) if v}
                cur = "esperando"
                continue
            if cur == "esperando" and isinstance(a, (int, float)):
                g = lambda n, d=0: (r[cols[n]] if n in cols and r[cols[n]] is not None else d)
                cur = {"nrc": int(a), "sec": str(g("SECC", "")), "idl": g("ID LIGA", None), "l": g("LIGA", None),
                       "cr": int(g("CRED")), "h": int(g("H")), "ht": int(g("HT")), "pp": int(g("PP")),
                       "cap": int(g("CAPA")), "reg": int(g("REGI")), "blq": str(g("BLOQUE", "NO")), "mod": mod, "m": []}
                secs.append(cur)
                continue
            if a == "PABE":
                continue
            if isinstance(cur, dict):
                v = [x for x in r if x is not None]
                if len(r) >= 6 and r[2] and str(r[2]).strip().endswith(","):
                    pab, aula, dia, hora, idd, doc = r[:6]
                else:
                    continue  # clases no presenciales (sin pabellon/aula): no se cargan
                if not pab or not aula:
                    continue
                try:
                    idd = int(idd)
                except (TypeError, ValueError):
                    pass  # ej. "STAFF": se conserva como texto
                cur["m"].append([str(pab), str(aula), str(dia).strip(" ,")[:3].upper(), str(hora).strip(), idd, str(doc).strip()])
        if cod and secs:
            cursos.append({"c": cod, "n": nombres.get(cod, ""), "s": secs})
    orden = {c: i for i, c in enumerate(nombres)}
    cursos.sort(key=lambda c: orden.get(c["c"], 10**6))
    return k, cursos, periodos, [c for c in nombres if c not in {x["c"] for x in cursos}]


def main():
    ap = argparse.ArgumentParser()
    raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    ap.add_argument("--entrada", default=os.path.join(raiz, "entrada"))
    ap.add_argument("--salida", default=os.path.join(raiz, "datos"))
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    prog_nombres = {}
    pm = os.path.join(a.entrada, "principal.xlsm")
    if os.path.exists(pm):
        for r in openpyxl.load_workbook(pm, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True):
            v = [x for x in r if x]
            if len(v) >= 2 and re.match(r"^[A-Z]{4}$", str(v[-2])):
                prog_nombres[str(v[-2])] = str(v[-1])
    D, avisos, periodos = [], [], set()
    archivos = sorted(glob.glob(os.path.join(a.entrada, "*.xlsm")))
    orden = list(prog_nombres)
    archivos.sort(key=lambda f: orden.index(os.path.basename(f)[:-5].upper()) if os.path.basename(f)[:-5].upper() in orden else 999)
    for f in archivos:
        if os.path.basename(f).lower() == "principal.xlsm":
            continue
        k, cursos, per, faltan = leer_programa(f)
        periodos |= per
        D.append({"k": k, "n": prog_nombres.get(k, k), "c": cursos})
        print("OK  %-5s %3d cursos  %4d secciones" % (k, len(cursos), sum(len(c["s"]) for c in cursos)))
        if faltan:
            avisos.append("%s: %d curso(s) figuran en la lista del programa pero no tienen hoja con secciones (ej. %s)" % (k, len(faltan), ", ".join(faltan[:4])))
    json.dump(D, open(os.path.join(a.salida, "horarios.json"), "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    esperados = [p for p in prog_nombres if p not in {x["k"] for x in D}]
    if esperados:
        avisos.append("Programas de principal.xlsm sin archivo .xlsm: " + ", ".join(esperados))
    json.dump({"generado": datetime.datetime.now().isoformat(timespec="seconds"), "periodo": sorted(periodos)[-1] if periodos else "",
               "avisos": avisos, "programas": len(D),
               "secciones": sum(len(c["s"]) for p in D for c in p["c"])},
              open(os.path.join(a.salida, "meta.json"), "w", encoding="utf-8"), ensure_ascii=False)
    for v in avisos:
        print("ADVERTENCIA:", v)
    print("Listo ->", a.salida, "| periodo:", sorted(periodos)[-1] if periodos else "?")


if __name__ == "__main__":
    main()
